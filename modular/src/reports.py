"""PDF and Excel report builders. A Report is a list of blocks that both formats can render."""
import io
import re
from dataclasses import dataclass, field
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
import streamlit as st
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .config import AMBER, APP_NAME, GREY, LINE, NAVY, TEAL


@dataclass
class Report:
    title: str
    meta: str
    blocks: list = field(default_factory=list)

    def heading(self, text):            self.blocks.append(("h", text))
    def kpis(self, items):              self.blocks.append(("kpis", items))
    def text(self, heading, lines):     self.blocks.append(("text", heading, lines))
    def bars(self, title, s, color=TEAL, fmt="{:,.0f}", unit=""): self.blocks.append(("bars", title, s, color, fmt, unit))
    def line(self, title, df):          self.blocks.append(("line", title, df))
    def table(self, title, df, note=""): self.blocks.append(("table", title, df, note))
    def extend(self, other):            self.blocks.extend(other.blocks)


def make_meta(b) -> str:
    r = b["report"]
    return (f"Data: {b['source']}, {r['rows']:,} rows from {r['customers']:,} customers, "
            f"{r['date_min']:%d %b %Y} to {r['date_max']:%d %b %Y}. Model: {st.session_state.get('model_name', 'Original model')}. "
            f"Generated {datetime.now():%d %b %Y %H:%M}.")


# ---------------------------------------------------------------- charts
def _png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=170, bbox_inches="tight", transparent=False, facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


def _bars_png(s, color, fmt, unit):
    s = s.dropna()
    fig, ax = plt.subplots(figsize=(7.2, 0.42 * max(len(s), 1) + 0.5))
    labels = [str(i).replace("_", " ") for i in s.index]
    ax.barh(labels, s.values, color=color, height=0.62)
    ax.invert_yaxis()
    top = float(s.max()) if len(s) and s.max() > 0 else 1.0
    for i, v in enumerate(s.values):
        ax.text(v + top * 0.01, i, fmt.format(v) + unit, va="center", fontsize=9, color=NAVY)
    ax.set_xlim(0, top * 1.22)
    ax.xaxis.set_visible(False)
    for sp in ("top", "right", "bottom"):
        ax.spines[sp].set_visible(False)
    ax.spines["left"].set_color(LINE)
    ax.tick_params(axis="y", length=0, labelsize=9, labelcolor=NAVY)
    return _png(fig)


def _line_png(df):
    fig, ax = plt.subplots(figsize=(7.2, 2.7))
    styles = [(NAVY, "-"), (TEAL, "--"), (AMBER, "-")]
    for (c, (col, ls)) in zip(df.columns, styles):
        ax.plot(df.index, df[c], color=col, linestyle=ls, linewidth=2, label=str(c))
    ax.grid(axis="y", color=LINE)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.tick_params(labelsize=8, labelcolor=GREY)
    ax.legend(frameon=False, fontsize=9, loc="upper center", bbox_to_anchor=(0.5, 1.14), ncol=3)
    fig.autofmt_xdate()
    return _png(fig)


# ---------------------------------------------------------------- PDF
def _styles():
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=20, textColor=colors.HexColor(NAVY),
                                alignment=0, spaceAfter=4),
        "meta": ParagraphStyle("m", parent=ss["Normal"], fontSize=8.5, textColor=colors.HexColor(GREY), leading=11),
        "h1": ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=16, textColor=colors.HexColor(TEAL),
                             spaceBefore=6, spaceAfter=6),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor(NAVY),
                             spaceBefore=10, spaceAfter=4),
        "body": ParagraphStyle("b", parent=ss["Normal"], fontSize=10, leading=14, textColor=colors.HexColor(NAVY)),
        "note": ParagraphStyle("n", parent=ss["Normal"], fontSize=8, textColor=colors.HexColor(GREY), leading=10),
        "k_l": ParagraphStyle("kl", parent=ss["Normal"], fontSize=8, textColor=colors.HexColor(GREY), leading=10),
        "k_v": ParagraphStyle("kv", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=13, textColor=colors.HexColor(NAVY), leading=16),
        "k_s": ParagraphStyle("ks", parent=ss["Normal"], fontSize=7.5, textColor=colors.HexColor(TEAL), leading=9),
        "cell": ParagraphStyle("c", parent=ss["Normal"], fontSize=7.5, leading=9, textColor=colors.HexColor(NAVY)),
        "cellh": ParagraphStyle("ch", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=7.5, leading=9, textColor=colors.white),
    }


def _esc(x):
    return str(x).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _kpi_table(items, st_):
    n = len(items)
    cells = [[Paragraph(_esc(l), st_["k_l"]), ] for l, v, s in items]
    row = [[Paragraph(_esc(l), st_["k_l"]), Paragraph(_esc(v), st_["k_v"]), Paragraph(_esc(s), st_["k_s"])] for l, v, s in items]
    t = Table([row], colWidths=[174 * mm / n] * n)
    t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(LINE)),
                           ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor(LINE)),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6),
                           ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    return t


def _data_table(df, st_, max_rows=25):
    d = df.head(max_rows)
    data = [[Paragraph(_esc(c), st_["cellh"]) for c in d.columns]]
    for row in d.itertuples(index=False):
        data.append([Paragraph(_esc(f"{v:,.2f}".rstrip("0").rstrip(".") if isinstance(v, float) else v), st_["cell"]) for v in row])
    t = Table(data, colWidths=[174 * mm / len(d.columns)] * len(d.columns), repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NAVY)),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F7")]),
                           ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor(LINE)),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    return t


def to_pdf(rep: Report) -> bytes:
    st_ = _styles()
    buf = io.BytesIO()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor(GREY))
        canvas.drawString(18 * mm, 10 * mm, f"{APP_NAME}. Synthetic data for demonstration.")
        canvas.drawRightString(192 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=rep.title, author=APP_NAME)
    story = [Paragraph(_esc(rep.title), st_["title"]), Paragraph(_esc(rep.meta), st_["meta"]), Spacer(1, 8)]
    for blk in rep.blocks:
        kind = blk[0]
        if kind == "h":
            story += [PageBreak() if len(story) > 4 else Spacer(1, 2), Paragraph(_esc(blk[1]), st_["h1"])]
        elif kind == "kpis":
            story += [_kpi_table(blk[1], st_), Spacer(1, 8)]
        elif kind == "text":
            if blk[1]:
                story.append(Paragraph(_esc(blk[1]), st_["h2"]))
            for ln in blk[2]:
                story.append(Paragraph("&bull; " + _esc(ln), st_["body"]))
            story.append(Spacer(1, 6))
        elif kind == "bars":
            _, title, s, color, fmt, unit = blk
            if len(s):
                png = _bars_png(s, color, fmt, unit)
                story += [Paragraph(_esc(title), st_["h2"]), Image(png, width=170 * mm, height=170 * mm * (0.42 * len(s) + 0.5) / 7.2)]
        elif kind == "line":
            if len(blk[2]):
                story += [Paragraph(_esc(blk[1]), st_["h2"]), Image(_line_png(blk[2]), width=170 * mm, height=170 * mm * 2.7 / 7.2)]
        elif kind == "table":
            _, title, df, note = blk
            story.append(Paragraph(_esc(title), st_["h2"]))
            if df is None or df.empty:
                story.append(Paragraph("No rows to show.", st_["note"]))
            else:
                story.append(_data_table(df, st_, 60))
                extra = f"Showing the first 60 of {len(df):,} rows. The Excel file has the full list. " if len(df) > 60 else ""
                if extra or note:
                    story.append(Paragraph(_esc(extra + note), st_["note"]))
            story.append(Spacer(1, 6))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ---------------------------------------------------------------- Excel
def _sheet_name(title, used):
    base = re.sub(r"\s*\([^)]*\)", "", title)           # drop units such as (GB)
    base = re.sub(r"[\[\]\*\?/\\:]", "", base)
    for a, b in (("Average data per customer per day", "Daily use"), ("Data per customer per day", "Daily use")):
        base = base.replace(a, b)
    base = (base[:28].rsplit(" ", 1)[0] if len(base) > 28 else base) or "Sheet"
    name, i = base, 2
    while name in used:
        name, i = f"{base[:25]} {i}", i + 1
    used.add(name)
    return name


def to_excel(rep: Report) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    used = {"Summary"}
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        wb_rows = [[rep.title], [rep.meta], []]
        for blk in rep.blocks:
            if blk[0] == "h":
                wb_rows += [[], [blk[1]]]
            elif blk[0] == "kpis":
                wb_rows += [["Measure", "Value", "Note"]] + [list(i) for i in blk[1]] + [[]]
            elif blk[0] == "text":
                wb_rows += ([[blk[1]]] if blk[1] else []) + [[ln] for ln in blk[2]] + [[]]
        pd.DataFrame(wb_rows).to_excel(xw, sheet_name="Summary", header=False, index=False)
        for blk in rep.blocks:
            if blk[0] == "bars":
                d = blk[2].rename_axis("Item").reset_index()
                d.columns = ["Item", "Value"]
                d["Item"] = d["Item"].astype(str).str.replace("_", " ")
                d.to_excel(xw, sheet_name=_sheet_name(blk[1], used), index=False)
            elif blk[0] == "line":
                blk[2].rename_axis("Date").reset_index().to_excel(xw, sheet_name=_sheet_name(blk[1], used), index=False)
            elif blk[0] == "table" and blk[2] is not None and not blk[2].empty:
                blk[2].to_excel(xw, sheet_name=_sheet_name(blk[1], used), index=False)
        for name, ws in xw.sheets.items():
            hdr = PatternFill("solid", fgColor=NAVY.lstrip("#"))
            if name == "Summary":
                ws["A1"].font = Font(name="Arial", bold=True, size=14)
                for row in ws.iter_rows(min_row=2):
                    for c in row:
                        c.font = Font(name="Arial", size=10)
                ws.column_dimensions["A"].width = 38
                ws.column_dimensions["B"].width = 22
                ws.column_dimensions["C"].width = 48
                continue
            for c in ws[1]:
                c.font = Font(name="Arial", bold=True, color="FFFFFF")
                c.fill = hdr
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 40)
            ws.freeze_panes = "A2"
    return buf.getvalue()


# ---------------------------------------------------------------- UI
def panel(key, build, title="Download report", extra=""):
    """An expander with a Prepare button. The report is built only on click, then offered as PDF and Excel."""
    ss = st.session_state
    sig = str([ss["version"], ss["model_name"], ss["filters"], ss["threshold"], ss["buffer"], ss["topup_days"],
               ss["topup_gb"], ss["price_gb"], ss["takeup"], ss["retain_days"], ss["retain_gb"],
               len(ss.get("imp_cache", {})), len(ss.get("shap_cache", {})), extra])
    store = ss.setdefault("report_store", {})
    with st.expander(title):
        if st.button("Prepare report", key=f"prep_{key}"):
            with st.spinner("Building the report"):
                rep = build()
                store[key] = (sig, to_pdf(rep), to_excel(rep), re.sub(r"\W+", "_", rep.title).strip("_").lower())
        item = store.get(key)
        if item and item[0] == sig:
            c1, c2 = st.columns(2)
            c1.download_button("Download PDF", item[1], f"{item[3]}.pdf", "application/pdf", key=f"pdf_{key}")
            c2.download_button("Download Excel", item[2], f"{item[3]}.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=f"xls_{key}")
        elif item:
            st.caption("The data or settings changed. Prepare the report again.")
        else:
            st.caption("Builds a PDF and an Excel file from what this page shows.")
